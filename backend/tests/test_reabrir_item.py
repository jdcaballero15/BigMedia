"""Reabrir un ítem enviado a facturar, avisar a Administración y reenviarlo (HU-067)."""

import io
import json
import smtplib
from datetime import date
from decimal import Decimal
from typing import Any

import pytest
from django.core import mail
from django.test import Client
from openpyxl import load_workbook

from apps.auditoria.models import RegistroAuditoria
from apps.clientes.models import Cliente
from apps.facturacion.models import ItemFacturacion

URL = "/api/facturacion/items/"
CORREO = "administracion@example.com"


@pytest.fixture
def con_correo(settings):
    settings.FACTURACION_CORREO_ADMINISTRACION = CORREO
    settings.DEFAULT_FROM_EMAIL = "avisos@example.com"


def _post(client: Client, url: str, data: dict[str, Any]) -> Any:
    return client.post(url, data=json.dumps(data), content_type="application/json")


def _patch(client: Client, item: ItemFacturacion, data: dict[str, Any]) -> Any:
    return client.patch(f"{URL}{item.pk}/", data=json.dumps(data), content_type="application/json")


def _reabrir(client: Client, item: ItemFacturacion, motivo: Any = "Cambió el neto") -> Any:
    return _post(client, f"{URL}{item.pk}/reabrir/", {"motivo": motivo})


def _enviado(**campos: Any) -> ItemFacturacion:
    datos = {
        "mes": date(2026, 11, 1),
        "cliente": Cliente.objects.create(nombre="Alfa SA"),
        "descripcion": "Alquiler noviembre",
        "neto": Decimal("10000.00"),
        "estado": "enviado",
        **campos,
    }
    return ItemFacturacion.objects.create(**datos)


def _enviar_mes(client: Client) -> Any:
    return _post(client, "/api/facturacion/envios/", {"mes": "2026-11"})


# --- Criterio 1: un Enviado no se edita sin reabrir ---


@pytest.mark.django_db
def test_un_item_enviado_no_se_edita_sin_reabrirlo(operativo_client):
    item = _enviado()
    respuesta = _patch(operativo_client, item, {"neto": "1.00"})
    assert respuesta.status_code == 400
    assert "estado" in respuesta.json()
    item.refresh_from_db()
    assert item.neto == Decimal("10000.00")


# --- Criterio 2: reabrir pide motivo y registra quién, cuándo y por qué ---


@pytest.mark.django_db
def test_reabrir_vuelve_a_borrador_y_registra_quien_cuando_y_por_que(operativo_client):
    item = _enviado()
    respuesta = _reabrir(operativo_client, item, " Cambió el neto ")
    assert respuesta.status_code == 200, respuesta.json()
    cuerpo = respuesta.json()
    assert cuerpo["estado"] == "borrador"
    assert (cuerpo["reabierto_por_nombre"], cuerpo["motivo_reapertura"]) == (
        "operativo",
        "Cambió el neto",
    )
    assert cuerpo["reabierto_en"] is not None
    assert cuerpo["reenviado"] is False
    registros = RegistroAuditoria.objects.filter(entidad="item_facturacion", entidad_id=item.pk)
    assert set(registros.values_list("campo", "valor_anterior", "valor_nuevo")) == {
        ("estado", "enviado", "borrador"),
        ("motivo_reapertura", "", "Cambió el neto"),
    }
    assert {r.usuario.username for r in registros if r.usuario} == {"operativo"}


@pytest.mark.django_db
@pytest.mark.parametrize("motivo", ["", "   ", None, 5])
def test_reabrir_sin_motivo_da_400(operativo_client, motivo):
    item = _enviado()
    respuesta = _reabrir(operativo_client, item, motivo)
    assert respuesta.status_code == 400
    assert "motivo" in respuesta.json()
    item.refresh_from_db()
    assert item.estado == "enviado"


@pytest.mark.django_db
@pytest.mark.parametrize(
    "campos",
    [
        {"estado": "borrador"},
        {"estado": "facturado", "numero_factura": "A-1"},
        {"estado": "anulado"},
    ],
)
def test_solo_se_reabre_un_item_enviado(operativo_client, campos):
    item = _enviado(**campos)
    respuesta = _reabrir(operativo_client, item)
    assert respuesta.status_code == 400
    assert "estado" in respuesta.json()


@pytest.mark.django_db
def test_el_lector_no_reabre(lector_client):
    assert _reabrir(lector_client, _enviado()).status_code == 403


@pytest.mark.django_db
def test_reabrir_un_item_inexistente_da_404(operativo_client):
    assert _post(operativo_client, f"{URL}999999/reabrir/", {"motivo": "x"}).status_code == 404


# --- Criterio 3: Administración recibe los avisos (RN-022) ---


@pytest.mark.django_db
@pytest.mark.regla("RN-022")
@pytest.mark.usefixtures("con_correo")
def test_reabrir_avisa_a_administracion(operativo_client, django_capture_on_commit_callbacks):
    item = _enviado()
    with django_capture_on_commit_callbacks(execute=True):
        _reabrir(operativo_client, item)
    assert len(mail.outbox) == 1
    aviso = mail.outbox[0]
    assert aviso.to == [CORREO]
    assert aviso.from_email == "avisos@example.com"
    assert f"#{item.pk}" in aviso.subject
    assert "Alquiler noviembre" in aviso.body
    assert "operativo" in aviso.body
    assert "Cambió el neto" in aviso.body
    assert "No lo facturen" in aviso.body


@pytest.mark.django_db
@pytest.mark.regla("RN-022")
@pytest.mark.usefixtures("con_correo")
def test_reenviar_avisa_con_los_valores_anteriores_y_nuevos(
    operativo_client, django_capture_on_commit_callbacks
):
    item = _enviado()
    with django_capture_on_commit_callbacks(execute=True):
        _reabrir(operativo_client, item)
        _patch(operativo_client, item, {"neto": "12000.00", "descripcion": "Alquiler y lona"})
        _post(operativo_client, f"{URL}{item.pk}/costos/", {"tipo": "otro", "monto": "500.00"})
        _enviar_mes(operativo_client)
    assert len(mail.outbox) == 2
    aviso = mail.outbox[1]
    assert "reenviado" in aviso.subject.lower()
    assert "Importe neto: 10000.00 → 12000.00" in aviso.body
    assert "Descripción: Alquiler noviembre → Alquiler y lona" in aviso.body
    assert "Costos: 0.00 → 500.00" in aviso.body
    assert "Neto final: 10000.00 → 11500.00" in aviso.body
    assert "Cliente" not in aviso.body


@pytest.mark.django_db
@pytest.mark.usefixtures("con_correo")
def test_reenviar_sin_cambios_lo_avisa(operativo_client, django_capture_on_commit_callbacks):
    item = _enviado()
    with django_capture_on_commit_callbacks(execute=True):
        _reabrir(operativo_client, item)
        _enviar_mes(operativo_client)
    assert "Sin cambios" in mail.outbox[1].body


@pytest.mark.django_db
@pytest.mark.usefixtures("con_correo")
def test_un_envio_comun_no_manda_avisos(operativo_client, django_capture_on_commit_callbacks):
    _enviado(estado="borrador")
    with django_capture_on_commit_callbacks(execute=True):
        assert _enviar_mes(operativo_client).status_code == 201
    assert mail.outbox == []


@pytest.mark.django_db
def test_sin_direccion_configurada_no_manda_correo_y_reabre_igual(
    operativo_client, settings, django_capture_on_commit_callbacks, caplog
):
    settings.FACTURACION_CORREO_ADMINISTRACION = ""
    item = _enviado()
    with django_capture_on_commit_callbacks(execute=True):
        respuesta = _reabrir(operativo_client, item)
    assert respuesta.status_code == 200
    assert mail.outbox == []
    assert "FACTURACION_CORREO_ADMINISTRACION" in caplog.text


@pytest.mark.django_db
@pytest.mark.usefixtures("con_correo")
def test_si_el_correo_falla_queda_en_el_log_y_reabre_igual(
    operativo_client, monkeypatch, django_capture_on_commit_callbacks, caplog
):
    def falla(*_args: Any, **_kwargs: Any) -> int:
        raise smtplib.SMTPException("sin conexión")

    monkeypatch.setattr("apps.facturacion.services.send_mail", falla)
    item = _enviado()
    with django_capture_on_commit_callbacks(execute=True):
        respuesta = _reabrir(operativo_client, item)
    assert respuesta.status_code == 200
    assert "No se pudo enviar el aviso" in caplog.text
    item.refresh_from_db()
    assert item.estado == "borrador"


@pytest.mark.django_db
@pytest.mark.usefixtures("con_correo")
def test_si_la_transaccion_falla_no_se_manda_correo(
    operativo_client, django_capture_on_commit_callbacks
):
    item = _enviado(estado="facturado", numero_factura="A-1")
    with django_capture_on_commit_callbacks(execute=True):
        _reabrir(operativo_client, item)
    assert mail.outbox == []


# --- Criterio 4: el reenvío vuelve a Enviado y queda marcado ---


@pytest.mark.django_db
def test_el_reenvio_vuelve_a_enviado_y_queda_marcado(operativo_client):
    item = _enviado()
    _reabrir(operativo_client, item)
    _patch(operativo_client, item, {"neto": "12000.00"})
    respuesta = _enviar_mes(operativo_client)
    assert respuesta.status_code == 201
    assert respuesta.json()["envio"]["items"] == [item.pk]
    detalle = operativo_client.get(f"{URL}{item.pk}/").json()
    assert (detalle["estado"], detalle["reenviado"]) == ("enviado", True)


@pytest.mark.django_db
@pytest.mark.regla("RN-024")
def test_la_planilla_marca_los_reenviados(operativo_client):
    item = _enviado()
    _enviado(estado="borrador", cliente=Cliente.objects.create(nombre="Beta SA"))
    _reabrir(operativo_client, item)
    envio = _enviar_mes(operativo_client).json()["envio"]["id"]
    respuesta = operativo_client.get(f"/api/facturacion/envios/{envio}/planilla/")
    hoja = load_workbook(io.BytesIO(respuesta.content)).active
    assert hoja is not None
    filas = [list(fila) for fila in hoja.iter_rows(values_only=True)]
    assert filas[0][-1] == "Reenviado"
    assert [(fila[1], fila[-1]) for fila in filas[1:]] == [("Alfa SA", "Sí"), ("Beta SA", None)]


@pytest.mark.django_db
def test_reabrir_otra_vez_un_reenviado_le_saca_la_marca_hasta_el_proximo_envio(operativo_client):
    item = _enviado()
    _reabrir(operativo_client, item)
    _enviar_mes(operativo_client)
    respuesta = _reabrir(operativo_client, item, "Otro ajuste")
    assert (respuesta.json()["reenviado"], respuesta.json()["motivo_reapertura"]) == (
        False,
        "Otro ajuste",
    )
