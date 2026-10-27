"""Planilla .xlsx del envío a facturar (HU-066, criterios 3 y 5; HU-062 y HU-064)."""

import io
import json
from datetime import date
from decimal import Decimal
from typing import Any

import pytest
from django.test import Client
from openpyxl import load_workbook

from apps.clientes.models import Cliente, Marca
from apps.facturacion.models import GrupoFactura, ItemFacturacion, Sector, TipoVenta

URL = "/api/facturacion/envios/"

COLUMNAS = [
    "Mes",
    "Cliente",
    "CUIT",
    "Marca",
    "Descripción",
    "Período desde",
    "Período hasta",
    "Cantidad",
    "Precio unitario",
    "Importe neto",
    "Alícuota IVA",
    "IVA",
    "Total",
    "Tipo de servicio",
    "Tipo de soporte",
    "Gob/Priv",
    "Tipo de venta",
    "Notas",
    "Orden de compra",
    "Fecha OC",
    "Grupo de factura",
]


def _enviar(client: Client) -> int:
    respuesta = client.post(
        URL, data=json.dumps({"mes": "2026-11"}), content_type="application/json"
    )
    assert respuesta.status_code == 201, respuesta.json()
    envio_id: int = respuesta.json()["envio"]["id"]
    return envio_id


def _planilla(client: Client, envio_id: int) -> tuple[Any, list[list[Any]]]:
    respuesta = client.get(f"{URL}{envio_id}/planilla/")
    assert respuesta.status_code == 200
    hoja = load_workbook(io.BytesIO(respuesta.content)).active
    assert hoja is not None
    return respuesta, [list(fila) for fila in hoja.iter_rows(values_only=True)]


@pytest.mark.django_db
@pytest.mark.regla("RN-024")
def test_planilla_con_las_columnas_y_los_importes_como_numeros(operativo_client):
    cliente = Cliente.objects.create(nombre="Zeta SA")
    marca = Marca.objects.create(nombre="Marca Z", cliente=cliente)
    ItemFacturacion.objects.create(
        mes=date(2026, 11, 1),
        cliente=cliente,
        cuit_facturacion="20123456786",
        marca=marca,
        descripcion="Alquiler noviembre",
        periodo_desde=date(2026, 11, 1),
        periodo_hasta=date(2026, 11, 30),
        cantidad=Decimal("10"),
        precio_unitario=Decimal("1000"),
        neto=Decimal("10000.00"),
        tipo_servicio="alquiler",
        tipo_soporte="Gran Formato",
        sector=Sector.objects.get(nombre="Privado"),
        tipo_venta=TipoVenta.objects.get(nombre="F"),
        notas="Facturar con la OC anterior",
    )
    respuesta, filas = _planilla(operativo_client, _enviar(operativo_client))
    assert respuesta["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert respuesta["Content-Disposition"].startswith('attachment; filename="facturacion-2026-11-')
    assert filas[0] == COLUMNAS
    fila = dict(zip(COLUMNAS, filas[1], strict=True))
    assert fila["Mes"] == "2026-11"
    assert (fila["Cliente"], fila["CUIT"], fila["Marca"]) == ("Zeta SA", "20123456786", "Marca Z")
    assert fila["Período desde"].date() == date(2026, 11, 1)
    assert (fila["Importe neto"], fila["IVA"], fila["Total"]) == (10000, 2100, 12100)
    assert isinstance(fila["Importe neto"], int | float)
    assert fila["Alícuota IVA"] == 21
    assert (fila["Tipo de servicio"], fila["Gob/Priv"], fila["Tipo de venta"]) == (
        "Alquiler",
        "Privado",
        "F",
    )
    assert (fila["Orden de compra"], fila["Fecha OC"], fila["Grupo de factura"]) == (
        None,
        None,
        None,
    )


@pytest.mark.django_db
def test_planilla_solo_con_los_items_del_envio_ordenados_por_cliente(operativo_client):
    for nombre in ("Zeta SA", "Alfa SA"):
        ItemFacturacion.objects.create(
            mes=date(2026, 11, 1),
            cliente=Cliente.objects.create(nombre=nombre),
            neto=Decimal("1.00"),
        )
    ItemFacturacion.objects.create(
        mes=date(2026, 11, 1), cliente=Cliente.objects.create(nombre="Anulado SA"), estado="anulado"
    )
    _, filas = _planilla(operativo_client, _enviar(operativo_client))
    assert [fila[1] for fila in filas[1:]] == ["Alfa SA", "Zeta SA"]


@pytest.mark.django_db
def test_planilla_de_envio_inexistente_da_404(lector_client):
    assert lector_client.get(f"{URL}999999/planilla/").status_code == 404


@pytest.mark.django_db
@pytest.mark.regla("RN-024")
def test_planilla_con_la_orden_de_compra_y_el_grupo(operativo_client):
    grupo = GrupoFactura.objects.create(mes=date(2026, 11, 1))
    ItemFacturacion.objects.create(
        mes=date(2026, 11, 1),
        cliente=Cliente.objects.create(nombre="Alfa SA"),
        neto=Decimal("1.00"),
        orden_compra_numero="OC-77",
        orden_compra_fecha=date(2026, 10, 20),
        grupo=grupo,
    )
    _, filas = _planilla(operativo_client, _enviar(operativo_client))
    fila = dict(zip(COLUMNAS, filas[1], strict=True))
    assert fila["Orden de compra"] == "OC-77"
    assert fila["Fecha OC"].date() == date(2026, 10, 20)
    assert fila["Grupo de factura"] == f"G-{grupo.pk}"


@pytest.mark.django_db
def test_los_items_de_un_grupo_salen_juntos(operativo_client):
    """HU-064 criterio 3 y HU-066 criterio 3: juntos y con el mismo identificador."""
    grupo = GrupoFactura.objects.create(mes=date(2026, 11, 1))
    alfa, beta, zeta = (Cliente.objects.create(nombre=n) for n in ("Alfa SA", "Beta SA", "Zeta SA"))
    for cliente, en_grupo in ((alfa, True), (beta, False), (zeta, True), (alfa, False)):
        ItemFacturacion.objects.create(
            mes=date(2026, 11, 1),
            cliente=cliente,
            neto=Decimal("1.00"),
            grupo=grupo if en_grupo else None,
        )
    _, filas = _planilla(operativo_client, _enviar(operativo_client))
    codigo = f"G-{grupo.pk}"
    assert [(fila[1], fila[-1]) for fila in filas[1:]] == [
        ("Alfa SA", codigo),
        ("Zeta SA", codigo),
        ("Alfa SA", None),
        ("Beta SA", None),
    ]


@pytest.mark.django_db
def test_planilla_con_el_cuit_del_item_distinto_del_cliente(operativo_client):
    """HU-064 criterio 4: el cliente no cambia y la planilla muestra el CUIT del ítem."""
    cliente = Cliente.objects.create(nombre="Alfa SA", cuit="20123456786")
    ItemFacturacion.objects.create(
        mes=date(2026, 11, 1), cliente=cliente, cuit_facturacion="27123456780", neto=Decimal("1")
    )
    _, filas = _planilla(operativo_client, _enviar(operativo_client))
    fila = dict(zip(COLUMNAS, filas[1], strict=True))
    assert fila["CUIT"] == "27123456780"
    cliente.refresh_from_db()
    assert cliente.cuit == "20123456786"
